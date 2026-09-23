# Copyright (c) 2026 Y2038.com LLC
# SPDX-License-Identifier: Apache-2.0

"""DOCX → plain text conversion."""

from __future__ import annotations

import io

from tads.ingest.archives import preflight_docx_package
from tads.ingest.fetch import IngestError
from tads.ingest.limits import assert_converted_text_limit, note_converted_chars
from tads.ingest.types import (
    DEFAULT_MAX_ARCHIVE_EXPANSION_RATIO,
    DEFAULT_MAX_ARCHIVE_MEMBER_BYTES,
    DEFAULT_MAX_CONTAINER_MEMBERS,
    DEFAULT_MAX_CONTAINER_UNCOMPRESSED_BYTES,
    DEFAULT_MAX_CONVERTED_CHARS,
)


def docx_to_text(
    data: bytes,
    *,
    max_converted_chars: int = DEFAULT_MAX_CONVERTED_CHARS,
    max_archive_member_bytes: int = DEFAULT_MAX_ARCHIVE_MEMBER_BYTES,
    max_archive_expansion_ratio: float = DEFAULT_MAX_ARCHIVE_EXPANSION_RATIO,
    max_container_members: int = DEFAULT_MAX_CONTAINER_MEMBERS,
    max_container_uncompressed_bytes: int = DEFAULT_MAX_CONTAINER_UNCOMPRESSED_BYTES,
) -> str:
    preflight_docx_package(
        data,
        max_archive_member_bytes=max_archive_member_bytes,
        max_archive_expansion_ratio=max_archive_expansion_ratio,
        max_container_members=max_container_members,
        max_container_uncompressed_bytes=max_container_uncompressed_bytes,
    )
    try:
        from docx import Document
    except ImportError as exc:
        raise IngestError(
            "python-docx is required for .docx conversion. "
            "Install with: pip install python-docx"
        ) from exc
    try:
        document = Document(io.BytesIO(data))
    except IngestError:
        raise
    except Exception as exc:  # noqa: BLE001 - surface corrupt docs cleanly
        raise IngestError(f"Failed to read DOCX: {exc}") from exc

    parts: list[str] = []
    used = 0
    for paragraph in document.paragraphs:
        used = note_converted_chars(paragraph.text, used=used, max_chars=max_converted_chars)
        parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ") for cell in row.cells
            ]
            if any(cells):
                line = " | ".join(cells)
                used = note_converted_chars(line, used=used, max_chars=max_converted_chars)
                parts.append(line)
    text = "\n".join(parts)
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    text = text.strip() + ("\n" if text.strip() else "")
    assert_converted_text_limit(text, max_converted_chars)
    return text